"""
core/parser.py
──────────────
PPT 結構解析：將每頁投影片轉換為結構化資料
"""

import re
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


# 用 top 位置（英吋）區分大標題 vs 副標題
TITLE_TOP_THRESHOLD = 0.6
EMU_PER_INCH = 914400


def _top_inch(shape) -> float:
    """將 shape.top（EMU）轉換為英吋"""
    return (shape.top or 0) / EMU_PER_INCH


def parse_slide(slide) -> dict:
    """
    解析單一投影片，回傳結構化資料：

    title       : 大標題文字（top < 0.6 英吋的「標題」shape）
    subtitle    : 副標題文字（top >= 0.6 英吋的「標題」shape）
    paragraphs  : 內容文字方塊的段落列表，每個元素為 (text, is_code)
                  is_code=True 代表疑似程式碼（以 # 開頭）
    shapes_text : 圖形內文字（橢圓圖說文字等），已依數字前綴排序
    table_rows  : 表格資料（list of list），無表格則為 []
    notes       : 備註欄文字
    """
    title = ""
    subtitle = ""
    content_shapes = []   # (top, paragraphs_list)
    shape_texts = []      # (order_key, text)
    table_rows = []

    for shape in slide.shapes:
        name   = shape.name
        s_type = shape.shape_type
        top    = _top_inch(shape)

        # ── 表格 ──────────────────────────────────────────
        if s_type == MSO_SHAPE_TYPE.TABLE:
            tbl = shape.table
            for row in tbl.rows:
                table_rows.append([
                    cell.text_frame.text.strip() for cell in row.cells
                ])
            continue

        if not shape.has_text_frame:
            continue

        # ── 大標題 / 副標題 ───────────────────────────────
        if "標題" in name and "圖說" not in name:
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if top < TITLE_TOP_THRESHOLD:
                title = text
            else:
                subtitle = text
            continue

        # ── 圖形內文字（橢圓圖說文字等）──────────────────
        if "圖說文字" in name:
            text = shape.text_frame.text.strip()
            if not text:
                continue
            m = re.match(r"^(\d+)[\.、]", text)
            order = int(m.group(1)) if m else 999
            shape_texts.append((order, text))
            continue

        # ── 一般內容文字方塊 ──────────────────────────────
        if "文字版面配置區" in name or "文字方塊" in name:
            # 保留段落結構（每個 paragraph 單獨存）
            paras = []
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                # 偵測疑似程式碼行（以 # 開頭）
                is_code = text.startswith("#")
                paras.append((text, is_code))
            if paras:
                content_shapes.append((top, paras))
            continue

    # 排序
    content_shapes.sort(key=lambda x: x[0])
    shape_texts.sort(key=lambda x: x[0])

    # 備註欄
    notes = ""
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()

    return {
        "title":       title,
        "subtitle":    subtitle,
        "paragraphs":  [p for _, paras in content_shapes for p in paras],
        "shapes_text": [t for _, t in shape_texts],
        "table_rows":  table_rows,
        "notes":       notes,
    }


def parse_all_slides(pptx_path: str) -> list:
    """解析整份 PPT，回傳 list of dict"""
    prs = Presentation(pptx_path)
    return [parse_slide(slide) for slide in prs.slides]
